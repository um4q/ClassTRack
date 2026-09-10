"""
Tests for app.excel_store.ExcelStore - the generic dataclass<->row (de)serializer,
typed CRUD wrappers, Settings key/value sheet, and empty-workbook bootstrap.

Uses the `store` fixture (private tmp copy of the real Fall 2026 workbook) for
realistic-data regressions, and `empty_store` (brand-new empty workbook) for
clean-slate edge cases. See tests/conftest.py for both.
"""
from __future__ import annotations

from datetime import date

import openpyxl
import pytest

from app import config
from app.excel_store import ExcelStore
from app.models import (
    Assessment, ChecklistItem, Course, PreLab, ScheduleRow, Topic,
)


# NOTE: ExcelStore.backup_now() (called by save() on every save where the
# workbook already exists - i.e. every reload round-trip test below
# triggers it) used to always resolve its destination from the global
# config.BACKUPS_DIR regardless of store.path, which would have copied
# every tmp workbook in this file straight into the REAL project's
# data/backups/ directory. Fixed to derive backups from store.path.parent
# (see ExcelStore.backups_dir) - store (a tmp copy) now backs up entirely
# within pytest's own tmp_path with no redirect needed. See
# test_backup_now_derives_destination_from_store_path_not_global_config
# below for a direct test of this.


def test_backup_now_derives_destination_from_store_path_not_global_config(store, monkeypatch, tmp_path):
    """Regression test for the design smell above: point config.BACKUPS_DIR
    somewhere backup_now() must NOT use, and confirm the backup still lands
    next to store.path instead."""
    decoy = tmp_path / "decoy_should_not_be_used"
    monkeypatch.setattr(config, "BACKUPS_DIR", decoy)

    dest = store.backup_now()

    assert dest is not None
    assert dest.parent == store.path.parent / "backups"
    assert dest.exists()
    assert not decoy.exists()


# --------------------------------------------------------------------------
# "=TRUE()" formula-as-boolean quirk
# --------------------------------------------------------------------------
def test_get_course_active_true_despite_excel_formula_string(store):
    """The real workbook stores Courses.active as the literal formula string
    "=TRUE()" (an Excel checkbox artifact), not a native bool cell. Confirm
    the raw cell really is that formula string, then confirm ExcelStore
    still decodes it to the native bool True."""
    wb = openpyxl.load_workbook(store.path, data_only=False)
    ws = wb["Courses"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    active_col = headers.index("active") + 1
    id_col = headers.index("course_id") + 1
    row1_active = next(
        ws.cell(row=r, column=active_col).value
        for r in range(2, ws.max_row + 1)
        if ws.cell(row=r, column=id_col).value == 1
    )
    assert isinstance(row1_active, str)
    assert row1_active.strip().upper() == "=TRUE()"

    course = store.get_course(1)
    assert course is not None
    assert course.active is True
    assert course in store.list_courses(active_only=True)


def test_course_active_bool_round_trips_natively_via_empty_store(empty_store):
    """A freshly-added course encodes/decodes `active` as a real bool cell
    (not a formula string) and round-trips correctly through save+reload,
    including the False case."""
    c = empty_store.add_course(Course(code="CMTC2341", name="Test Course", active=False))
    empty_store.save(force=True)
    empty_store.reload()

    reloaded = empty_store.get_course(c.course_id)
    assert reloaded is not None
    assert reloaded.active is False
    assert reloaded not in empty_store.list_courses(active_only=True)


# --------------------------------------------------------------------------
# Optional[float] clearing regression
# --------------------------------------------------------------------------
def test_assessment_score_set_then_cleared_round_trips_through_reload(store):
    """Real regression: `ws.cell(row, col, value=None)` is a silent no-op in
    openpyxl, so clearing an Optional[float] back to None used to leave the
    old value on disk. Assessment #1 in the real workbook starts with
    score=None; set it, persist+reload, then clear it back to None,
    persist+reload again, and confirm it is actually None - not still 88.5."""
    a = store.get_assessment(1)
    assert a is not None
    assert a.score is None  # sanity check on the real fixture data

    a.score = 88.5
    store.update_assessment(a)
    store.save(force=True)
    store.reload()

    reloaded = store.get_assessment(1)
    assert reloaded.score == pytest.approx(88.5)

    reloaded.score = None
    store.update_assessment(reloaded)
    store.save(force=True)
    store.reload()

    cleared = store.get_assessment(1)
    assert cleared.score is None


# --------------------------------------------------------------------------
# Delete clears the whole row (same underlying openpyxl gotcha)
# --------------------------------------------------------------------------
def test_delete_topic_blanks_every_cell_not_just_filters_it_out(empty_store, qapp):
    t = empty_store.add_topic(Topic(
        course_id=1, unit="U1", section="1.1", title="Serial comms",
        status="Learning", confidence=4, priority=2, link="foo.pdf",
        page=12, estimated_hours=2.5, notes="delete-me regression row",
    ))
    empty_store.save(force=True)
    empty_store.reload()
    assert any(x.topic_id == t.topic_id for x in empty_store.list_topics())

    assert empty_store.delete_topic(t.topic_id) is True
    empty_store.save(force=True)
    empty_store.reload()

    remaining = empty_store.list_topics()
    assert all(x.topic_id != t.topic_id for x in remaining)

    # Confirm via a completely independent ExcelStore instance too, not
    # just the one that performed the delete.
    fresh = ExcelStore(empty_store.path)
    fresh.load()
    assert all(x.topic_id != t.topic_id for x in fresh.list_topics())

    # And read the raw bytes directly: every cell in that row must be
    # genuinely blank (None), not merely absent from list_topics()'s output
    # because the app's own reader skips rows with a blank anchor column.
    wb = openpyxl.load_workbook(empty_store.path, data_only=False)
    ws = wb["Topics"]
    row_idx = 2  # the only row ever written in this sheet
    for col in range(1, ws.max_column + 1):
        assert ws.cell(row=row_idx, column=col).value is None


# --------------------------------------------------------------------------
# Checklist round-trip (list[ChecklistItem], pipe-delimited)
# --------------------------------------------------------------------------
def test_prelab_checklist_round_trips_mixed_checked_state_and_order(empty_store):
    checklist = [
        ChecklistItem(text="Read the lab section in CP#1400", checked=True),
        ChecklistItem(text="Bring memory stick", checked=False),
        ChecklistItem(text="Answer pre-lab questions", checked=True),
        ChecklistItem(text="Print the lab sheet", checked=False),
    ]
    p = empty_store.add_prelab(PreLab(
        course_id=1, lab_number="X02-01", title="Lab 1", checklist=checklist,
    ))
    empty_store.save(force=True)
    empty_store.reload()

    reloaded = empty_store.get_prelab(p.prelab_id)
    assert reloaded is not None
    assert [(i.text, i.checked) for i in reloaded.checklist] == [
        (i.text, i.checked) for i in checklist
    ]


# --------------------------------------------------------------------------
# List-field round-trips: List[int] and List[date]
# --------------------------------------------------------------------------
def test_assessment_topic_ids_list_int_round_trip(empty_store):
    ids = [5, 12, 1, 100]
    a = empty_store.add_assessment(Assessment(course_id=1, title="Midterm", topic_ids=ids))
    empty_store.save(force=True)
    empty_store.reload()

    reloaded = empty_store.get_assessment(a.assessment_id)
    assert reloaded is not None
    assert reloaded.topic_ids == ids  # order preserved, not just membership


def test_assessment_topic_ids_empty_list_round_trips_to_empty_not_none(empty_store):
    a = empty_store.add_assessment(Assessment(course_id=1, title="No topics", topic_ids=[]))
    empty_store.save(force=True)
    empty_store.reload()

    reloaded = empty_store.get_assessment(a.assessment_id)
    assert reloaded.topic_ids == []


def test_schedule_row_dates_list_date_round_trip_biweekly_friday_lab(empty_store):
    """The real workbook's X02 lab pattern: a Friday lab section that meets
    only on specific bi-weekly dates rather than every week."""
    biweekly_fridays = [
        date(2026, 9, 4), date(2026, 9, 18), date(2026, 10, 2),
        date(2026, 10, 16), date(2026, 10, 30),
    ]
    s = empty_store.add_schedule_row(ScheduleRow(
        course_id=1, kind="Lab", section="X02", day_of_week="Friday",
        dates=biweekly_fridays,
    ))
    empty_store.save(force=True)
    empty_store.reload()

    reloaded = empty_store.get_schedule_row(s.schedule_id)
    assert reloaded is not None
    assert reloaded.dates == biweekly_fridays


# --------------------------------------------------------------------------
# Lenient date parsing from raw (hand-edited-in-Excel-style) strings
# --------------------------------------------------------------------------
def test_topic_next_review_parses_lenient_raw_date_string_formats(empty_store):
    topic_ids = [
        empty_store.add_topic(Topic(course_id=1, unit="U", title=f"Topic {i}")).topic_id
        for i in range(4)
    ]
    empty_store.save(force=True)

    # Bypass ExcelStore entirely - write raw strings straight into the cell,
    # simulating a student hand-editing the sheet in Excel.
    raw_wb = openpyxl.load_workbook(empty_store.path, data_only=False)
    ws = raw_wb["Topics"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    id_col = headers.index("topic_id") + 1
    nr_col = headers.index("next_review") + 1

    raw_by_id = {
        topic_ids[0]: "2026-9-3",           # ISO-ish, unpadded month/day
        topic_ids[1]: "09/03/2026",         # explicit M/D/Y format
        topic_ids[2]: "September 3, 2026",  # only the dateutil fallback parses this
        topic_ids[3]: "not-a-real-date",    # unparsable
    }
    for r in range(2, ws.max_row + 1):
        tid = ws.cell(row=r, column=id_col).value
        if tid in raw_by_id:
            ws.cell(row=r, column=nr_col).value = raw_by_id[tid]
    raw_wb.save(empty_store.path)

    empty_store.reload()
    topics = {t.topic_id: t for t in empty_store.list_topics()}

    expected = date(2026, 9, 3)
    assert topics[topic_ids[0]].next_review == expected
    assert topics[topic_ids[1]].next_review == expected
    assert topics[topic_ids[2]].next_review == expected
    # Documented fallback for a genuinely unparsable string: None, not a crash.
    assert topics[topic_ids[3]].next_review is None


# --------------------------------------------------------------------------
# Settings: get/set/date/int/bool round-trip + missing-key fallback
# --------------------------------------------------------------------------
def test_get_setting_and_set_setting_round_trip_through_reload(empty_store):
    assert empty_store.get_setting("theme") == config.SETTINGS_DEFAULTS["theme"]

    empty_store.set_setting("theme", "light")
    assert empty_store.get_setting("theme") == "light"

    empty_store.save(force=True)
    empty_store.reload()
    assert empty_store.get_setting("theme") == "light"


def test_get_setting_missing_key_falls_back_to_settings_defaults(empty_store):
    # Simulate an older workbook whose Settings sheet is missing a key that
    # a newer config.SETTINGS_DEFAULTS introduced, by blanking that row
    # directly (bypassing ExcelStore).
    wb = openpyxl.load_workbook(empty_store.path, data_only=False)
    ws = wb["Settings"]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "theme":
            ws.cell(row=row, column=1).value = None
            ws.cell(row=row, column=2).value = None
            break
    wb.save(empty_store.path)

    empty_store.reload()
    assert empty_store.get_setting("theme") == config.SETTINGS_DEFAULTS["theme"]

    # A key that has never existed anywhere at all.
    assert empty_store.get_setting("totally_unknown_key") is None
    assert empty_store.get_setting("totally_unknown_key", "fallback") == "fallback"


def test_setting_date_int_bool_helpers_and_their_fallbacks(empty_store):
    assert empty_store.setting_date("semester_start") == date(2026, 9, 1)
    assert empty_store.setting_int("default_study_block_min") == 50
    assert empty_store.setting_bool("notifications") is True

    # Invalid int value falls back to the given default rather than raising.
    empty_store.set_setting("default_study_block_min", "not-a-number")
    assert empty_store.setting_int("default_study_block_min", default=99) == 99

    # Blank bool value falls back to the given default.
    empty_store.set_setting("notifications", "")
    assert empty_store.setting_bool("notifications", default=False) is False

    # A date setting that was never set at all.
    assert empty_store.setting_date("nonexistent_date_key") is None


# --------------------------------------------------------------------------
# Empty-workbook bootstrap contract
# --------------------------------------------------------------------------
def test_empty_store_creates_every_sheet_with_correct_headers_and_settings(empty_store):
    assert empty_store.created_empty is True

    wb = openpyxl.load_workbook(empty_store.path, data_only=False)
    assert set(wb.sheetnames) == set(config.SHEET_HEADERS.keys())

    for sheet, headers in config.SHEET_HEADERS.items():
        if not headers:
            continue  # Overview/README are formulas/free-text sheets
        ws = wb[sheet]
        actual = [ws.cell(row=1, column=i + 1).value for i in range(len(headers))]
        assert actual == headers, f"{sheet} header row mismatch: {actual!r} != {headers!r}"

    assert empty_store.all_settings() == dict(config.SETTINGS_DEFAULTS)
