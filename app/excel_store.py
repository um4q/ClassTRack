"""
ExcelStore - the ONLY module that touches the workbook.

Owns: loading/creating data/study_tracker.xlsx, a generic dataclass<->row
(de)serializer driven by models.py's type hints, typed CRUD wrappers per
sheet, the Settings key/value sheet, autosave debouncing, timestamped
backups, and graceful handling of the file being locked open in Excel.

Every other module reads/writes workbook data exclusively through a shared
ExcelStore instance - never via openpyxl directly.
"""
from __future__ import annotations

import dataclasses
import logging
import re
import shutil
import typing
from datetime import date, datetime, time
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from PySide6.QtCore import QObject, QTimer, Signal

from app import config, models
from app.models import (
    AttendanceEntry, BrainDumpNote, ChecklistItem, Course, Goal, GradeWeight,
    Holiday, Material, MockExam, PracticeQuestion, PreLab, RoadmapItem,
    ScheduleRow, StudyLogEntry, SyllabusFlag, Topic, Assessment,
)

# --------------------------------------------------------------------------
# Logging - rotating file handler shared by the whole app (import this
# module first; every other module does `logger = logging.getLogger(
# "study_tracker")` to reuse the same handler).
# --------------------------------------------------------------------------
LOGGER_NAME = "study_tracker"


def _setup_logging() -> logging.Logger:
    log = logging.getLogger(LOGGER_NAME)
    if log.handlers:
        return log
    log.setLevel(logging.INFO)
    try:
        config.DATA_DIR.mkdir(parents=True, exist_ok=True)
        handler = RotatingFileHandler(
            config.LOG_PATH, maxBytes=1_000_000, backupCount=3, encoding="utf-8"
        )
    except OSError:
        handler = logging.StreamHandler()
    handler.setFormatter(
        logging.Formatter("%(asctime)s %(levelname)s %(name)s: %(message)s")
    )
    log.addHandler(handler)
    return log


logger = _setup_logging()

EMPTY_README_LINES = [
    "This workbook was auto-created because data/study_tracker.xlsx was missing.",
    "",
    "Restore the supplied study_tracker.xlsx for your real Fall 2026 semester",
    "data (courses, syllabus topics, exam dates, X02 lab sessions, etc.).",
    "",
    "You can also bulk-edit any sheet directly in Excel - Study Tracker reloads",
    "on launch, or press Ctrl+R to reload without restarting.",
    "",
    "Sheets: Courses, Materials, GradeWeights, Topics, Assessments, PreLabs,",
    "Schedule, Holidays, SyllabusFlags, StudyLog, Attendance, Goals,",
    "PracticeQuestions, MockExams, Roadmap, BrainDump, Settings.",
]


class ExcelStore(QObject):
    """Shared, in-memory-cached view of data/study_tracker.xlsx."""

    dataChanged = Signal(str)   # sheet name, or "" after a full load/reload
    saved = Signal()
    saveError = Signal(str)
    lockWarning = Signal(bool)  # True = workbook is locked open elsewhere

    def __init__(self, path: Optional[Path] = None, parent=None):
        super().__init__(parent)
        self.path: Path = Path(path) if path else config.WORKBOOK_PATH
        self._wb: Optional[Workbook] = None
        self._dirty = False
        self._locked = False
        self._settings: dict[str, str] = {}
        self.created_empty = False
        # Cache of sheet title -> {header name: column index}, since a
        # header row never changes for the life of a loaded workbook -
        # invalidated on every load()/reload() (profiled: ~2.6% of a full
        # click-through session was spent re-scanning row 1 on every single
        # list_*/add_*/update_*/delete_* call before this cache existed).
        self._header_cache: dict[str, dict[str, int]] = {}

        self._autosave_timer = QTimer(self)
        self._autosave_timer.setSingleShot(True)
        self._autosave_timer.timeout.connect(lambda: self.save())

        self._lock_retry_timer = QTimer(self)
        self._lock_retry_timer.setSingleShot(False)
        self._lock_retry_timer.timeout.connect(lambda: self.save())

    # ---------------------------------------------------------- lifecycle --
    def load(self) -> None:
        """Open the workbook, creating an empty one (with headers/README) if
        it doesn't exist or fails to parse. Sets ``created_empty`` so the UI
        can show the "empty workbook created" banner (see settings/dashboard)."""
        self.created_empty = False
        self._header_cache = {}
        if self.path.exists():
            try:
                self._wb = openpyxl.load_workbook(self.path, data_only=False)
            except Exception:
                logger.exception("Failed to open %s - creating an empty workbook", self.path)
                self._wb = self._create_empty_workbook()
                self.created_empty = True
        else:
            logger.info("No workbook at %s - creating an empty one", self.path)
            self.path.parent.mkdir(parents=True, exist_ok=True)
            self._wb = self._create_empty_workbook()
            self.created_empty = True

        self._ensure_all_sheets()
        self._load_settings()
        self._dirty = False
        if self.created_empty:
            self.save(force=True)
        self.dataChanged.emit("")

    def reload(self) -> None:
        """Re-read from disk, discarding any unsaved in-memory edits."""
        self.load()

    def save(self, force: bool = False) -> bool:
        if not self._dirty and not force:
            return True
        if self._wb is None:
            return False
        try:
            if self.path.exists():
                self.backup_now()
            self.path.parent.mkdir(parents=True, exist_ok=True)
            self._wb.save(self.path)
            self._dirty = False
            if self._locked:
                self._locked = False
                self._lock_retry_timer.stop()
                self.lockWarning.emit(False)
            self.saved.emit()
            return True
        except PermissionError as e:
            logger.warning("Workbook is locked (open in Excel?): %s", e)
            if not self._locked:
                self._locked = True
                self.lockWarning.emit(True)
            if not self._lock_retry_timer.isActive():
                self._lock_retry_timer.start(config.SAVE_LOCK_RETRY_MS)
            self.saveError.emit(str(e))
            return False
        except Exception as e:
            logger.exception("Failed to save workbook")
            self.saveError.emit(str(e))
            return False

    def close(self) -> bool:
        """Called from the main window's closeEvent - flush any pending edit."""
        self._autosave_timer.stop()
        return self.save(force=True)

    def mark_dirty(self) -> None:
        self._dirty = True
        self._autosave_timer.start(config.AUTOSAVE_DEBOUNCE_MS)

    def is_dirty(self) -> bool:
        return self._dirty

    def backup_now(self) -> Optional[Path]:
        if not self.path.exists():
            return None
        try:
            config.BACKUPS_DIR.mkdir(parents=True, exist_ok=True)
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            dest = config.BACKUPS_DIR / f"study_tracker_{ts}.xlsx"
            shutil.copy2(self.path, dest)
            self._prune_backups()
            return dest
        except OSError:
            logger.exception("Backup failed")
            return None

    def _prune_backups(self) -> None:
        backups = sorted(
            config.BACKUPS_DIR.glob("study_tracker_*.xlsx"),
            key=lambda p: p.stat().st_mtime,
        )
        while len(backups) > config.BACKUP_KEEP_COUNT:
            oldest = backups.pop(0)
            oldest.unlink(missing_ok=True)

    def resolve_path(self, rel: Optional[str]) -> Optional[Path]:
        """Resolve a workbook-relative path (Materials.path, Topics.link, ...)
        against data/ - relative paths in the workbook are relative to data/."""
        if not rel:
            return None
        # Strip only a TRAILING "#page=N" suffix - filenames in this export
        # legitimately contain a literal "#" (e.g. "CP#1400 - CMTC 2341.pdf"),
        # so splitting on the first "#" anywhere (the old behavior) truncated
        # every CMTC2341 link to "...Content/CP", breaking every Open button
        # for that course's materials/topics/prelabs.
        rel = re.sub(r"#page=\d+$", "", rel, flags=re.IGNORECASE)
        p = Path(rel)
        return p if p.is_absolute() else (config.DATA_DIR / rel)

    # ------------------------------------------------------------ Settings --
    def _load_settings(self) -> None:
        settings = dict(config.SETTINGS_DEFAULTS)
        ws = self._wb["Settings"] if "Settings" in self._wb.sheetnames else None
        if ws is not None:
            header_map = self._header_map(ws)
            kcol, vcol = header_map.get("key"), header_map.get("value")
            if kcol and vcol:
                for row_idx in range(2, ws.max_row + 1):
                    k = ws.cell(row=row_idx, column=kcol).value
                    if not k:
                        continue
                    v = ws.cell(row=row_idx, column=vcol).value
                    settings[str(k).strip()] = "" if v is None else str(v)
        self._settings = settings

    def get_setting(self, key: str, default=None):
        if key in self._settings:
            return self._settings[key]
        return default if default is not None else config.SETTINGS_DEFAULTS.get(key)

    def all_settings(self) -> dict[str, str]:
        return dict(self._settings)

    def set_setting(self, key: str, value) -> None:
        value_str = "" if value is None else str(value)
        self._settings[key] = value_str
        ws = self._wb["Settings"]
        header_map = self._header_map(ws)
        kcol, vcol = header_map.get("key"), header_map.get("value")
        row_idx = None
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=kcol).value == key:
                row_idx = r
                break
        if row_idx is None:
            row_idx = ws.max_row + 1
            ws.cell(row=row_idx, column=kcol, value=key)
        ws.cell(row=row_idx, column=vcol, value=value_str)
        self.mark_dirty()
        self.dataChanged.emit("Settings")

    def setting_date(self, key: str) -> Optional[date]:
        return self._parse_date(self.get_setting(key))

    def setting_int(self, key: str, default: int = 0) -> int:
        try:
            return int(float(self.get_setting(key, default)))
        except (TypeError, ValueError):
            return default

    def setting_bool(self, key: str, default: bool = False) -> bool:
        v = self.get_setting(key)
        return self._decode_bool(v) if v not in (None, "") else default

    # ------------------------------------------------- generic CRUD engine --
    def _header_map(self, ws) -> dict[str, int]:
        # Returned dict is the cached instance itself (not a copy) for
        # speed - every caller in this file only reads it (.get()); never
        # mutate a returned header map, or every other caller sharing this
        # sheet's cache entry would see the corruption too.
        cached = self._header_cache.get(ws.title)
        if cached is not None:
            return cached
        m: dict[str, int] = {}
        for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=1):
            if cell.value:
                m[str(cell.value).strip()] = idx
        self._header_cache[ws.title] = m
        return m

    def _default_of(self, f: dataclasses.Field):
        if f.default is not dataclasses.MISSING:
            return f.default
        if f.default_factory is not dataclasses.MISSING:  # type: ignore[misc]
            return f.default_factory()
        return None

    def _decode_bool(self, raw) -> bool:
        if isinstance(raw, bool):
            return raw
        if raw is None:
            return False
        # Excel sometimes stores a checkbox-style boolean as a formula
        # string ("=TRUE()"/"=FALSE()") rather than a literal - strip that.
        s = str(raw).strip().lower().strip("=()")
        return s in ("true", "1", "yes", "y")

    def _parse_date(self, raw) -> Optional[date]:
        if raw in (None, ""):
            return None
        if isinstance(raw, datetime):
            return raw.date()
        if isinstance(raw, date):
            return raw
        s = str(raw).strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
        try:
            from dateutil import parser as _dtparser
            return _dtparser.parse(s).date()
        except Exception:
            return None

    def _parse_time(self, raw) -> Optional[time]:
        if raw in (None, ""):
            return None
        if isinstance(raw, datetime):
            return raw.time().replace(second=0, microsecond=0)
        if isinstance(raw, time):
            return raw
        s = str(raw).strip()
        for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p"):
            try:
                return datetime.strptime(s, fmt).time()
            except ValueError:
                continue
        return None

    def _parse_value(self, f: dataclasses.Field, resolved_type, raw, delim: str):
        origin = typing.get_origin(resolved_type)
        if origin is typing.Union:
            inner = [a for a in typing.get_args(resolved_type) if a is not type(None)][0]
            if raw in (None, ""):
                return None
            return self._parse_value(f, inner, raw, delim)
        if origin in (list,):
            elem_type = typing.get_args(resolved_type)[0]
            if raw in (None, ""):
                return []
            parts = [p for p in str(raw).split(delim)]
            if elem_type is ChecklistItem:
                items = []
                for p in parts:
                    p = p.strip()
                    if not p:
                        continue
                    checked = p[:3].lower() in ("[x]",)
                    text = p[3:].strip() if p[:1] == "[" and len(p) >= 3 else p
                    items.append(ChecklistItem(text=text, checked=checked))
                return items
            return [
                self._parse_value(f, elem_type, p.strip(), delim)
                for p in parts if p.strip() != ""
            ]
        if resolved_type is bool:
            return self._decode_bool(raw) if raw not in (None, "") else self._default_of(f)
        if resolved_type is int:
            if raw in (None, ""):
                return self._default_of(f)
            try:
                return int(float(raw))
            except (TypeError, ValueError):
                return self._default_of(f)
        if resolved_type is float:
            if raw in (None, ""):
                return self._default_of(f)
            try:
                return float(raw)
            except (TypeError, ValueError):
                return self._default_of(f)
        if resolved_type is date:
            return self._parse_date(raw) or self._default_of(f)
        if resolved_type is time:
            return self._parse_time(raw) or self._default_of(f)
        return "" if raw is None else str(raw)

    def _encode_scalar(self, t, value):
        if value is None:
            return None
        if t is bool:
            return bool(value)
        if t is date:
            return value.isoformat() if isinstance(value, date) else str(value)
        if t is time:
            return value.strftime("%H:%M") if isinstance(value, time) else str(value)
        if t in (int, float):
            return value
        return str(value)

    def _encode_value(self, resolved_type, value, delim: str):
        origin = typing.get_origin(resolved_type)
        if origin is typing.Union:
            inner = [a for a in typing.get_args(resolved_type) if a is not type(None)][0]
            return None if value is None else self._encode_value(inner, value, delim)
        if origin in (list,):
            elem_type = typing.get_args(resolved_type)[0]
            if not value:
                return None
            if elem_type is ChecklistItem:
                return delim.join(
                    f"[{'x' if it.checked else ' '}] {it.text}" for it in value
                )
            return delim.join(str(self._encode_scalar(elem_type, v)) for v in value)
        return self._encode_scalar(resolved_type, value)

    def _list_rows(self, sheet: str, cls: type) -> list:
        if self._wb is None or sheet not in self._wb.sheetnames:
            return []
        ws = self._wb[sheet]
        header_map = self._header_map(ws)
        hints = typing.get_type_hints(cls)
        flds = dataclasses.fields(cls)
        id_field = models.SHEET_ID_FIELD.get(sheet)
        anchor_field = id_field or flds[0].name
        anchor_col = header_map.get(anchor_field)
        out = []
        if not anchor_col or ws.max_row < 2:
            return out
        for row_idx in range(2, ws.max_row + 1):
            anchor_val = ws.cell(row=row_idx, column=anchor_col).value
            if anchor_val in (None, ""):
                continue
            try:
                kwargs = {}
                for f in flds:
                    col = header_map.get(f.name)
                    raw = ws.cell(row=row_idx, column=col).value if col else None
                    delim = f.metadata.get("delim", ",")
                    kwargs[f.name] = self._parse_value(f, hints[f.name], raw, delim)
                out.append(cls(**kwargs))
            except Exception:
                logger.warning("Skipping malformed row %d in %s", row_idx, sheet, exc_info=True)
        return out

    def _find_row(self, sheet: str, id_field: str, id_value) -> Optional[int]:
        ws = self._wb[sheet]
        header_map = self._header_map(ws)
        col = header_map.get(id_field)
        if not col:
            return None
        for row_idx in range(2, ws.max_row + 1):
            cv = ws.cell(row=row_idx, column=col).value
            if cv in (None, ""):
                continue
            try:
                if int(float(cv)) == int(id_value):
                    return row_idx
            except (TypeError, ValueError):
                continue
        return None

    def _find_blank_row(self, sheet: str, id_field: Optional[str]) -> int:
        ws = self._wb[sheet]
        header_map = self._header_map(ws)
        flds_first = next(iter(config.SHEET_HEADERS[sheet]), None)
        anchor = id_field or flds_first
        col = header_map.get(anchor, 1) if anchor else 1
        for row_idx in range(2, max(ws.max_row, 1) + 1):
            if ws.cell(row=row_idx, column=col).value in (None, ""):
                return row_idx
        return ws.max_row + 1

    def _write_row(self, sheet: str, cls: type, row_idx: int, obj) -> None:
        ws = self._wb[sheet]
        header_map = self._header_map(ws)
        hints = typing.get_type_hints(cls)
        for f in dataclasses.fields(cls):
            col = header_map.get(f.name)
            if not col:
                continue
            value = getattr(obj, f.name)
            delim = f.metadata.get("delim", ",")
            # NOTE: ws.cell(row, col, value=None) is a no-op in openpyxl, so
            # an Optional field being cleared back to blank must go through
            # .value explicitly or the old cell content would survive.
            ws.cell(row=row_idx, column=col).value = self._encode_value(hints[f.name], value, delim)

    def _next_id(self, sheet: str, id_field: str) -> int:
        ws = self._wb[sheet]
        header_map = self._header_map(ws)
        col = header_map.get(id_field)
        max_id = 0
        if col:
            for row_idx in range(2, ws.max_row + 1):
                v = ws.cell(row=row_idx, column=col).value
                try:
                    max_id = max(max_id, int(float(v)))
                except (TypeError, ValueError):
                    continue
        return max_id + 1

    def _add_row(self, sheet: str, cls: type, obj):
        id_field = models.SHEET_ID_FIELD.get(sheet)
        if id_field and not getattr(obj, id_field, 0):
            setattr(obj, id_field, self._next_id(sheet, id_field))
        row_idx = self._find_blank_row(sheet, id_field)
        self._write_row(sheet, cls, row_idx, obj)
        self.mark_dirty()
        self.dataChanged.emit(sheet)
        return obj

    def _update_row(self, sheet: str, cls: type, id_field: str, obj) -> None:
        row_idx = self._find_row(sheet, id_field, getattr(obj, id_field))
        if row_idx is None:
            raise ValueError(f"{sheet} row with {id_field}={getattr(obj, id_field)} not found")
        self._write_row(sheet, cls, row_idx, obj)
        self.mark_dirty()
        self.dataChanged.emit(sheet)

    def _delete_row(self, sheet: str, id_field: str, id_value) -> bool:
        row_idx = self._find_row(sheet, id_field, id_value)
        if row_idx is None:
            return False
        ws = self._wb[sheet]
        for col in range(1, ws.max_column + 1):
            # NOTE: ws.cell(row, col, value=None) is a no-op in openpyxl
            # (it only assigns when value is not None) - clear via .value.
            ws.cell(row=row_idx, column=col).value = None
        self.mark_dirty()
        self.dataChanged.emit(sheet)
        return True

    # -------------------------------------------------- empty-workbook build --
    def _apply_data_validations(self, ws, sheet: str, headers: list[str]) -> None:
        for col_idx, h in enumerate(headers, start=1):
            choices = config.DATA_VALIDATIONS.get((sheet, h))
            if not choices:
                continue
            dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
            col_letter = get_column_letter(col_idx)
            dv.add(f"{col_letter}2:{col_letter}1000")
            ws.add_data_validation(dv)

    def _create_empty_workbook(self) -> Workbook:
        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("Overview")
        ws["A1"] = "Study Tracker - empty workbook (no semester data loaded)"
        ws = wb.create_sheet("README")
        for i, line in enumerate(EMPTY_README_LINES, start=1):
            ws.cell(row=i, column=1, value=line)
        for sheet, headers in config.SHEET_HEADERS.items():
            if sheet in ("Overview", "README") or not headers:
                continue
            ws = wb.create_sheet(sheet)
            for col_idx, h in enumerate(headers, start=1):
                ws.cell(row=1, column=col_idx, value=h)
            self._apply_data_validations(ws, sheet, headers)
        ws = wb["Settings"]
        for i, (k, v) in enumerate(config.SETTINGS_DEFAULTS.items(), start=2):
            ws.cell(row=i, column=1, value=k)
            ws.cell(row=i, column=2, value=v)
        return wb

    def _ensure_all_sheets(self) -> None:
        for sheet, headers in config.SHEET_HEADERS.items():
            if sheet not in self._wb.sheetnames:
                ws = self._wb.create_sheet(sheet)
                for col_idx, h in enumerate(headers, start=1):
                    ws.cell(row=1, column=col_idx, value=h)
                if headers:
                    self._apply_data_validations(ws, sheet, headers)
                logger.warning("Sheet %s was missing - created with headers", sheet)

    # ====================================================================
    # Typed per-sheet wrappers
    # ====================================================================

    # -- Courses ----------------------------------------------------------
    def list_courses(self, active_only: bool = False) -> list[Course]:
        rows = self._list_rows("Courses", Course)
        return [c for c in rows if c.active] if active_only else rows

    def get_course(self, course_id: int) -> Optional[Course]:
        return next((c for c in self.list_courses() if c.course_id == course_id), None)

    def add_course(self, course: Course) -> Course:
        return self._add_row("Courses", Course, course)

    def update_course(self, course: Course) -> None:
        self._update_row("Courses", Course, "course_id", course)

    def delete_course(self, course_id: int) -> bool:
        return self._delete_row("Courses", "course_id", course_id)

    # -- Materials ----------------------------------------------------------
    def list_materials(self, course_id: Optional[int] = None) -> list[Material]:
        rows = self._list_rows("Materials", Material)
        rows.sort(key=lambda m: m.sort_order)
        return [m for m in rows if m.course_id == course_id] if course_id is not None else rows

    def get_material(self, material_id: int) -> Optional[Material]:
        return next((m for m in self.list_materials() if m.material_id == material_id), None)

    def add_material(self, material: Material) -> Material:
        return self._add_row("Materials", Material, material)

    def update_material(self, material: Material) -> None:
        self._update_row("Materials", Material, "material_id", material)

    def delete_material(self, material_id: int) -> bool:
        return self._delete_row("Materials", "material_id", material_id)

    # -- GradeWeights ---------------------------------------------------
    def list_grade_weights(self, course_id: Optional[int] = None) -> list[GradeWeight]:
        rows = self._list_rows("GradeWeights", GradeWeight)
        return [g for g in rows if g.course_id == course_id] if course_id is not None else rows

    def get_grade_weight(self, weight_id: int) -> Optional[GradeWeight]:
        return next((g for g in self.list_grade_weights() if g.weight_id == weight_id), None)

    def add_grade_weight(self, gw: GradeWeight) -> GradeWeight:
        return self._add_row("GradeWeights", GradeWeight, gw)

    def update_grade_weight(self, gw: GradeWeight) -> None:
        self._update_row("GradeWeights", GradeWeight, "weight_id", gw)

    def delete_grade_weight(self, weight_id: int) -> bool:
        return self._delete_row("GradeWeights", "weight_id", weight_id)

    # -- Topics -----------------------------------------------------------
    def list_topics(self, course_id: Optional[int] = None) -> list[Topic]:
        rows = self._list_rows("Topics", Topic)
        return [t for t in rows if t.course_id == course_id] if course_id is not None else rows

    def get_topic(self, topic_id: int) -> Optional[Topic]:
        return next((t for t in self.list_topics() if t.topic_id == topic_id), None)

    def add_topic(self, topic: Topic) -> Topic:
        return self._add_row("Topics", Topic, topic)

    def update_topic(self, topic: Topic) -> None:
        self._update_row("Topics", Topic, "topic_id", topic)

    def delete_topic(self, topic_id: int) -> bool:
        return self._delete_row("Topics", "topic_id", topic_id)

    # -- Assessments ------------------------------------------------------
    def list_assessments(self, course_id: Optional[int] = None) -> list[Assessment]:
        rows = self._list_rows("Assessments", Assessment)
        rows.sort(key=lambda a: (a.due_date is None, a.due_date or date.max))
        return [a for a in rows if a.course_id == course_id] if course_id is not None else rows

    def get_assessment(self, assessment_id: int) -> Optional[Assessment]:
        return next((a for a in self.list_assessments() if a.assessment_id == assessment_id), None)

    def add_assessment(self, a: Assessment) -> Assessment:
        return self._add_row("Assessments", Assessment, a)

    def update_assessment(self, a: Assessment) -> None:
        self._update_row("Assessments", Assessment, "assessment_id", a)

    def delete_assessment(self, assessment_id: int) -> bool:
        return self._delete_row("Assessments", "assessment_id", assessment_id)

    # -- PreLabs ------------------------------------------------------------
    def list_prelabs(self, course_id: Optional[int] = None) -> list[PreLab]:
        rows = self._list_rows("PreLabs", PreLab)
        rows.sort(key=lambda p: (p.lab_date is None, p.lab_date or date.max))
        return [p for p in rows if p.course_id == course_id] if course_id is not None else rows

    def get_prelab(self, prelab_id: int) -> Optional[PreLab]:
        return next((p for p in self.list_prelabs() if p.prelab_id == prelab_id), None)

    def add_prelab(self, p: PreLab) -> PreLab:
        return self._add_row("PreLabs", PreLab, p)

    def update_prelab(self, p: PreLab) -> None:
        self._update_row("PreLabs", PreLab, "prelab_id", p)

    def delete_prelab(self, prelab_id: int) -> bool:
        return self._delete_row("PreLabs", "prelab_id", prelab_id)

    # -- Schedule -----------------------------------------------------------
    def list_schedule(self, course_id: Optional[int] = None) -> list[ScheduleRow]:
        rows = self._list_rows("Schedule", ScheduleRow)
        return [s for s in rows if s.course_id == course_id] if course_id is not None else rows

    def get_schedule_row(self, schedule_id: int) -> Optional[ScheduleRow]:
        return next((s for s in self.list_schedule() if s.schedule_id == schedule_id), None)

    def add_schedule_row(self, s: ScheduleRow) -> ScheduleRow:
        return self._add_row("Schedule", ScheduleRow, s)

    def update_schedule_row(self, s: ScheduleRow) -> None:
        self._update_row("Schedule", ScheduleRow, "schedule_id", s)

    def delete_schedule_row(self, schedule_id: int) -> bool:
        return self._delete_row("Schedule", "schedule_id", schedule_id)

    # -- Holidays (small, effectively read-only from the app) --------------
    def list_holidays(self) -> list[Holiday]:
        rows = self._list_rows("Holidays", Holiday)
        rows.sort(key=lambda h: h.date or date.max)
        return rows

    def add_holiday(self, h: Holiday) -> Holiday:
        row_idx = self._find_blank_row("Holidays", None)
        self._write_row("Holidays", Holiday, row_idx, h)
        self.mark_dirty()
        self.dataChanged.emit("Holidays")
        return h

    # -- SyllabusFlags --------------------------------------------------
    def list_syllabus_flags(self, course_id: Optional[int] = None, unresolved_only: bool = False) -> list[SyllabusFlag]:
        rows = self._list_rows("SyllabusFlags", SyllabusFlag)
        if course_id is not None:
            rows = [f for f in rows if f.course_id == course_id]
        if unresolved_only:
            rows = [f for f in rows if not f.resolved]
        return rows

    def get_syllabus_flag(self, flag_id: int) -> Optional[SyllabusFlag]:
        return next((f for f in self.list_syllabus_flags() if f.flag_id == flag_id), None)

    def add_syllabus_flag(self, f: SyllabusFlag) -> SyllabusFlag:
        return self._add_row("SyllabusFlags", SyllabusFlag, f)

    def update_syllabus_flag(self, f: SyllabusFlag) -> None:
        self._update_row("SyllabusFlags", SyllabusFlag, "flag_id", f)

    def delete_syllabus_flag(self, flag_id: int) -> bool:
        return self._delete_row("SyllabusFlags", "flag_id", flag_id)

    # -- StudyLog -------------------------------------------------------
    def list_study_log(self, course_id: Optional[int] = None) -> list[StudyLogEntry]:
        rows = self._list_rows("StudyLog", StudyLogEntry)
        rows.sort(key=lambda s: (s.date is None, s.date or date.max))
        return [s for s in rows if s.course_id == course_id] if course_id is not None else rows

    def get_study_log_entry(self, log_id: int) -> Optional[StudyLogEntry]:
        return next((s for s in self.list_study_log() if s.log_id == log_id), None)

    def add_study_log_entry(self, s: StudyLogEntry) -> StudyLogEntry:
        return self._add_row("StudyLog", StudyLogEntry, s)

    def update_study_log_entry(self, s: StudyLogEntry) -> None:
        self._update_row("StudyLog", StudyLogEntry, "log_id", s)

    def delete_study_log_entry(self, log_id: int) -> bool:
        return self._delete_row("StudyLog", "log_id", log_id)

    # -- Attendance -------------------------------------------------------
    def list_attendance(self, course_id: Optional[int] = None) -> list[AttendanceEntry]:
        rows = self._list_rows("Attendance", AttendanceEntry)
        return [a for a in rows if a.course_id == course_id] if course_id is not None else rows

    def get_attendance_entry(self, attendance_id: int) -> Optional[AttendanceEntry]:
        return next((a for a in self.list_attendance() if a.attendance_id == attendance_id), None)

    def add_attendance_entry(self, a: AttendanceEntry) -> AttendanceEntry:
        return self._add_row("Attendance", AttendanceEntry, a)

    def update_attendance_entry(self, a: AttendanceEntry) -> None:
        self._update_row("Attendance", AttendanceEntry, "attendance_id", a)

    def delete_attendance_entry(self, attendance_id: int) -> bool:
        return self._delete_row("Attendance", "attendance_id", attendance_id)

    # -- Goals --------------------------------------------------------------
    def list_goals(self) -> list[Goal]:
        return self._list_rows("Goals", Goal)

    def get_goal(self, goal_id: int) -> Optional[Goal]:
        return next((g for g in self.list_goals() if g.goal_id == goal_id), None)

    def add_goal(self, g: Goal) -> Goal:
        return self._add_row("Goals", Goal, g)

    def update_goal(self, g: Goal) -> None:
        self._update_row("Goals", Goal, "goal_id", g)

    def delete_goal(self, goal_id: int) -> bool:
        return self._delete_row("Goals", "goal_id", goal_id)

    # -- PracticeQuestions ------------------------------------------------
    def list_practice_questions(self, course_id: Optional[int] = None) -> list[PracticeQuestion]:
        rows = self._list_rows("PracticeQuestions", PracticeQuestion)
        return [q for q in rows if q.course_id == course_id] if course_id is not None else rows

    def get_practice_question(self, question_id: int) -> Optional[PracticeQuestion]:
        return next((q for q in self.list_practice_questions() if q.question_id == question_id), None)

    def add_practice_question(self, q: PracticeQuestion) -> PracticeQuestion:
        return self._add_row("PracticeQuestions", PracticeQuestion, q)

    def update_practice_question(self, q: PracticeQuestion) -> None:
        self._update_row("PracticeQuestions", PracticeQuestion, "question_id", q)

    def delete_practice_question(self, question_id: int) -> bool:
        return self._delete_row("PracticeQuestions", "question_id", question_id)

    # -- MockExams ----------------------------------------------------------
    def list_mock_exams(self, course_id: Optional[int] = None) -> list[MockExam]:
        rows = self._list_rows("MockExams", MockExam)
        rows.sort(key=lambda m: (m.date_taken is None, m.date_taken or date.max))
        return [m for m in rows if m.course_id == course_id] if course_id is not None else rows

    def get_mock_exam(self, exam_id: int) -> Optional[MockExam]:
        return next((m for m in self.list_mock_exams() if m.exam_id == exam_id), None)

    def add_mock_exam(self, m: MockExam) -> MockExam:
        return self._add_row("MockExams", MockExam, m)

    def update_mock_exam(self, m: MockExam) -> None:
        self._update_row("MockExams", MockExam, "exam_id", m)

    def delete_mock_exam(self, exam_id: int) -> bool:
        return self._delete_row("MockExams", "exam_id", exam_id)

    # -- Roadmap --------------------------------------------------------
    def list_roadmap(self, course_id: Optional[int] = None) -> list[RoadmapItem]:
        rows = self._list_rows("Roadmap", RoadmapItem)
        rows.sort(key=lambda r: r.sort_order)
        return [r for r in rows if r.course_id == course_id] if course_id is not None else rows

    def get_roadmap_item(self, roadmap_id: int) -> Optional[RoadmapItem]:
        return next((r for r in self.list_roadmap() if r.roadmap_id == roadmap_id), None)

    def add_roadmap_item(self, r: RoadmapItem) -> RoadmapItem:
        return self._add_row("Roadmap", RoadmapItem, r)

    def update_roadmap_item(self, r: RoadmapItem) -> None:
        self._update_row("Roadmap", RoadmapItem, "roadmap_id", r)

    def delete_roadmap_item(self, roadmap_id: int) -> bool:
        return self._delete_row("Roadmap", "roadmap_id", roadmap_id)

    def delete_auto_roadmap(self, course_id: int) -> int:
        """Remove every auto_generated=True roadmap row for a course, so
        'Generate roadmap' can replace only what it previously wrote."""
        n = 0
        for r in self.list_roadmap(course_id):
            if r.auto_generated and self._delete_row("Roadmap", "roadmap_id", r.roadmap_id):
                n += 1
        return n

    # -- BrainDump ------------------------------------------------------
    def list_brain_dump(self, on_date: Optional[date] = None, course_id: Optional[int] = None) -> list[BrainDumpNote]:
        rows = self._list_rows("BrainDump", BrainDumpNote)
        rows.sort(key=lambda n: n.date or date.min, reverse=True)
        if on_date is not None:
            rows = [n for n in rows if n.date == on_date]
        if course_id is not None:
            rows = [n for n in rows if n.course_id == course_id]
        return rows

    def get_brain_dump_note(self, note_id: int) -> Optional[BrainDumpNote]:
        return next((n for n in self.list_brain_dump() if n.note_id == note_id), None)

    def add_brain_dump_note(self, n: BrainDumpNote) -> BrainDumpNote:
        return self._add_row("BrainDump", BrainDumpNote, n)

    def update_brain_dump_note(self, n: BrainDumpNote) -> None:
        self._update_row("BrainDump", BrainDumpNote, "note_id", n)

    def delete_brain_dump_note(self, note_id: int) -> bool:
        return self._delete_row("BrainDump", "note_id", note_id)
